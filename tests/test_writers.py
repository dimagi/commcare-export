import csv
import datetime
import io
import tempfile
import zipfile

import openpyxl
import sqlalchemy

import pytest
from commcare_export.specs import TableSpec
from commcare_export.writers import (
    CsvTableWriter,
    Excel2007TableWriter,
    JValueTableWriter,
    SqlTableWriter,
)


@pytest.fixture()
def writer(db_params):
    return SqlTableWriter(db_params['url'], poolclass=sqlalchemy.pool.NullPool)


@pytest.fixture()
def strict_writer(db_params):
    return SqlTableWriter(
        db_params['url'],
        poolclass=sqlalchemy.pool.NullPool,
        strict_types=True
    )


TYPE_MAP = {
    'mysql': {
        bool: lambda x: int(x)
    },
}


class TestWriters(object):

    def test_JValueTableWriter(self):
        writer = JValueTableWriter()
        writer.write_table(
            TableSpec(
                **{
                    'name':
                        'foo',
                    'headings': ['a', 'bjørn', 'c', 'd'],
                    'rows': [
                        [1, '2', 3, datetime.date(2015, 1, 1)],
                        [4, '日本', 6, datetime.date(2015, 1, 2)],
                    ]
                }
            )
        )

        writer.write_table(
            TableSpec(
                **{
                    'name': 'foo',
                    'headings': ['a', 'bjørn', 'c', 'd'],
                    'rows': [[5, 'bob', 9,
                              datetime.date(2018, 1, 2)],]
                }
            )
        )

        assert writer.tables == {
            'foo':
                TableSpec(
                    **{
                        'name':
                            'foo',
                        'headings': ['a', 'bjørn', 'c', 'd'],
                        'rows': [
                            [1, '2', 3, '2015-01-01'],
                            [4, '日本', 6, '2015-01-02'],
                            [5, 'bob', 9, '2018-01-02'],
                        ],
                    }
                )
        }

    def test_Excel2007TableWriter(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx') as file:
            with Excel2007TableWriter(file=file) as writer:
                writer.write_table(
                    TableSpec(
                        **{
                            'name': 'foo',
                            'headings': ['a', 'bjørn', 'c'],
                            'rows': [
                                [1, '2', 3],
                                [4, '日本', 6],
                            ]
                        }
                    )
                )

            self._check_Excel2007TableWriter_output(file.name)

    def test_Excel2007TableWriter_write_mutli(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx') as file:
            with Excel2007TableWriter(file=file) as writer:
                writer.write_table(
                    TableSpec(
                        **{
                            'name': 'foo',
                            'headings': ['a', 'bjørn', 'c'],
                            'rows': [[1, '2', 3],]
                        }
                    )
                )

                writer.write_table(
                    TableSpec(
                        **{
                            'name': 'foo',
                            'headings': ['a', 'bjørn', 'c'],
                            'rows': [[4, '日本', 6],]
                        }
                    )
                )
            self._check_Excel2007TableWriter_output(file.name)

    def _check_Excel2007TableWriter_output(self, filename):
        output_wb = openpyxl.load_workbook(filename)

        assert output_wb.sheetnames == ['foo']
        foo_sheet = output_wb['foo']
        assert [
            [cell.value for cell in row] for row in foo_sheet['A1:C3']
        ] == [
            ['a', 'bjørn', 'c'],
            ['1', '2', '3'
            ],  # Note how pyxl does some best-effort parsing to *whatever* type
            ['4', '日本', '6'],
        ]

    def test_CsvTableWriter(self):
        with tempfile.NamedTemporaryFile() as file:
            with CsvTableWriter(file=file) as writer:
                writer.write_table(
                    TableSpec(
                        **{
                            'name': 'foo',
                            'headings': ['a', 'bjørn', 'c'],
                            'rows': [
                                [1, '2', 3],
                                [4, '日本', 6],
                            ]
                        }
                    )
                )

            with zipfile.ZipFile(file.name, 'r') as output_zip:
                with output_zip.open('foo.csv') as csv_file:
                    output = csv.reader(
                        io.TextIOWrapper(csv_file, encoding='utf-8')
                    )

                    assert [row for row in output] == [
                        ['a', 'bjørn', 'c'],
                        ['1', '2', '3'],
                        ['4', '日本', '6'],
                    ]


@pytest.mark.dbtest
class TestSQLWriters(object):

    def _type_convert(self, connection, row):
        """
        Different databases store and return values differently so
        convert the values in the expected row to match the DB.
        """

        def convert(type_map, value):
            func = type_map.get(value.__class__, None)
            return func(value) if func else value

        for driver, type_map in TYPE_MAP.items():
            if driver in connection.engine.driver:
                return {k: convert(type_map, v) for k, v in row.items()}

        return row

    def test_insert(self, writer):
        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_insert',
                        'headings': ['id', 'a', 'b', 'c'],
                        'rows': [
                            ['bizzle', 1, 2, 3],
                            ['bazzle', 4, 5, 6],
                        ]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            result = dict([(row['id'], row) for row in writer.connection
                           .execute('SELECT id, a, b, c FROM foo_insert')])

        assert len(result) == 2
        assert dict(result['bizzle']) == {
            'id': 'bizzle',
            'a': 1,
            'b': 2,
            'c': 3
        }
        assert dict(result['bazzle']) == {
            'id': 'bazzle',
            'a': 4,
            'b': 5,
            'c': 6
        }

    def test_upsert(self, writer):
        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_upsert',
                        'headings': ['id', 'a', 'b', 'c'],
                        'rows': [['zing', 3, None, 5]]
                    }
                )
            )

        # don't select column 'b' since it hasn't been created yet
        with writer:
            result = dict([
                (row['id'], row) for row in
                writer.connection.execute('SELECT id, a, c FROM foo_upsert')
            ])
        assert len(result) == 1
        assert dict(result['zing']) == {'id': 'zing', 'a': 3, 'c': 5}

        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            'foo_upsert',
                        'headings': ['id', 'a', 'b', 'c'],
                        'rows': [
                            ['bizzle', 1, 'yo', 3],
                            ['bazzle', 4, '日本', 6],
                        ]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            result = dict([(row['id'], row) for row in writer.connection
                           .execute('SELECT id, a, b, c FROM foo_upsert')])

        assert len(result) == 3
        assert dict(result['bizzle']) == {
            'id': 'bizzle',
            'a': 1,
            'b': 'yo',
            'c': 3
        }
        assert dict(result['bazzle']) == {
            'id': 'bazzle',
            'a': 4,
            'b': '日本',
            'c': 6
        }

        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_upsert',
                        'headings': ['id', 'a', 'b', 'c'],
                        'rows': [['bizzle', 7, '本', 9],]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            result = dict([(row['id'], row) for row in writer.connection
                           .execute('SELECT id, a, b, c FROM foo_upsert')])

        assert len(result) == 3
        assert dict(result['bizzle']) == {
            'id': 'bizzle',
            'a': 7,
            'b': '本',
            'c': 9
        }
        assert dict(result['bazzle']) == {
            'id': 'bazzle',
            'a': 4,
            'b': '日本',
            'c': 6
        }

    def test_types(self, writer):
        self._test_types(writer, 'foo_fancy_types')

    def _test_types(self, writer, table_name):
        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            table_name,
                        'headings': ['id', 'a', 'b', 'c', 'd', 'e'],
                        'rows': [
                            [
                                'bizzle', 1, 'yo', True,
                                datetime.date(2015, 1, 1),
                                datetime.datetime(2014, 4, 2, 18, 56, 12)
                            ],
                            [
                                'bazzle', 4, '日本', False,
                                datetime.date(2015, 1, 2),
                                datetime.datetime(2014, 5, 1, 11, 16, 45)
                            ],
                        ]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            connection = writer.connection
            result = dict([
                (row['id'], row) for row in connection
                .execute('SELECT id, a, b, c, d, e FROM %s' % table_name)
            ])

            assert len(result) == 2
            expected = {
                'bizzle': {
                    'id': 'bizzle',
                    'a': 1,
                    'b': 'yo',
                    'c': True,
                    'd': datetime.date(2015, 1, 1),
                    'e': datetime.datetime(2014, 4, 2, 18, 56, 12)
                },
                'bazzle': {
                    'id': 'bazzle',
                    'a': 4,
                    'b': '日本',
                    'c': False,
                    'd': datetime.date(2015, 1, 2),
                    'e': datetime.datetime(2014, 5, 1, 11, 16, 45)
                }
            }

            for id, row in result.items():
                assert id in expected
                assert dict(row
                           ) == self._type_convert(connection, expected[id])

    def test_change_type(self, writer):
        self._test_types(writer, 'foo_fancy_type_changes')

        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            'foo_fancy_type_changes',
                        'headings': ['id', 'a', 'b', 'c', 'd', 'e'],
                        'rows': [[
                            'bizzle', 'yo dude', '本', 'true',
                            datetime.datetime(2015, 2, 13),
                            '2014-08-01T11:23:45:00.0000Z'
                        ],]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            result = dict([
                (row['id'], row) for row in writer.connection.execute(
                    'SELECT id, a, b, c, d, e FROM foo_fancy_type_changes'
                )
            ])

        assert len(result) == 2
        expected = {
            'bizzle': {
                'id': 'bizzle',
                'a': 'yo dude',
                'b': '本',
                'c': 'true',
                'd': datetime.date(2015, 2, 13),
                'e': '2014-08-01T11:23:45:00.0000Z'
            },
            'bazzle': {
                'id': 'bazzle',
                'a': '4',
                'b': '日本',
                'c': 'false',
                'd': datetime.date(2015, 1, 2),
                'e': '2014-05-01 11:16:45'
            }
        }

        if 'mysql' in writer.connection.engine.driver:
            # mysql weirdness
            expected['bazzle']['c'] = '0'
        if 'pyodbc' in writer.connection.engine.driver:
            expected['bazzle']['c'] = '0'
            # MSSQL includes fractional seconds in returned value.
            expected['bazzle']['e'] = '2014-05-01 11:16:45.0000000'

        for id, row in result.items():
            assert id in expected
            assert dict(row) == expected[id]

    def test_json_type(self, writer):
        complex_object = {
            'poke1': {
                'name': 'snorlax',
                'color': 'blue',
                'attributes': {
                    'strength': 10,
                    'endurance': 10,
                    'speed': 4,
                },
                'friends': [
                    'pikachu',
                    'charmander',
                ],
            },
            'poke2': {
                'name': 'pikachu',
                'color': 'yellow',
                'attributes': {
                    'strength': 2,
                    'endurance': 2,
                    'speed': 8,
                    'cuteness': 10,
                },
                'friends': [
                    'snorlax',
                    'charmander',
                ],
            },
        }
        with writer:
            if not writer.is_postgres:
                return
            writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_with_json',
                        'headings': ['id', 'json_col'],
                        'rows': [
                            ['simple', {
                                'k1': 'v1',
                                'k2': 'v2'
                            }],
                            ['with_lists', {
                                'l1': ['i1', 'i2']
                            }],
                            ['complex', complex_object],
                        ],
                        'data_types': [
                            'text',
                            'json',
                        ]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with writer:
            result = dict([(row['id'], row) for row in writer.connection
                           .execute('SELECT id, json_col FROM foo_with_json')])

        assert len(result) == 3
        assert dict(result['simple']) == {
            'id': 'simple',
            'json_col': {
                'k1': 'v1',
                'k2': 'v2'
            }
        }
        assert dict(result['with_lists']) == {
            'id': 'with_lists',
            'json_col': {
                'l1': ['i1', 'i2']
            }
        }
        assert dict(result['complex']) == {
            'id': 'complex',
            'json_col': complex_object
        }

    def test_explicit_types(self, strict_writer):
        with strict_writer:
            strict_writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_explicit_types',
                        'headings': ['id', 'a', 'b', 'c', 'd'],
                        'rows': [
                            ['bizzle', '1', 2, 3, '7'],
                            ['bazzle', '4', 5, 6, '8'],
                        ],
                        'data_types': [
                            'text',
                            'integer',
                            'text',
                            None,
                        ]
                    }
                )
            )

        # We can use raw SQL instead of SqlAlchemy expressions because
        # we built the DB above
        with strict_writer:
            result = dict([
                (row['id'], row) for row in strict_writer.connection
                .execute('SELECT id, a, b, c, d FROM foo_explicit_types')
            ])

        assert len(result) == 2
        # a casts strings to ints, b casts ints to text, c default falls back to ints, d default falls back to text
        assert dict(result['bizzle']) == {
            'id': 'bizzle',
            'a': 1,
            'b': '2',
            'c': 3,
            'd': '7'
        }
        assert dict(result['bazzle']) == {
            'id': 'bazzle',
            'a': 4,
            'b': '5',
            'c': 6,
            'd': '8'
        }

    def test_mssql_nvarchar_length_upsize(self, writer):
        with writer:
            if 'odbc' not in writer.connection.engine.driver:
                return

            # Initialize a table with columns where we expect the
            # "some_data" column to be of length 900 bytes, and the
            # "big_data" column to be of nvarchar(max)
            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            'mssql_nvarchar_length',
                        'headings': ['id', 'some_data', 'big_data'],
                        'rows': [
                            [
                                'bizzle', (b'\0' * 800).decode('utf-8'),
                                (b'\0' * 901).decode('utf-8')
                            ],
                            [
                                'bazzle', (b'\0' * 500).decode('utf-8'),
                                (b'\0' * 800).decode('utf-8')
                            ],
                        ]
                    }
                )
            )

            connection = writer.connection

            result = self._get_column_lengths(
                connection, 'mssql_nvarchar_length'
            )
            assert result['some_data'] == ('some_data', 'nvarchar', 900)
            # nvarchar(max) is listed as -1
            assert result['big_data'] == ('big_data', 'nvarchar', -1)

            # put bigger data into "some_column" to ensure it is resized
            # properly
            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            'mssql_nvarchar_length',
                        'headings': ['id', 'some_data', 'big_data'],
                        'rows': [[
                            'sizzle', (b'\0' * 901).decode('utf-8'),
                            (b'\0' * 901).decode('utf-8')
                        ],]
                    }
                )
            )

            result = self._get_column_lengths(
                connection, 'mssql_nvarchar_length'
            )
            assert result['some_data'] == ('some_data', 'nvarchar', -1)
            assert result['big_data'] == ('big_data', 'nvarchar', -1)

    def test_mssql_nvarchar_length_downsize(self, writer):
        with writer:
            if 'odbc' not in writer.connection.engine.driver:
                return

            # Initialize a table with NVARCHAR(max), and make sure
            # smaller data doesn't reduce the size of the column
            metadata = sqlalchemy.MetaData()
            create_sql = sqlalchemy.schema.CreateTable(
                sqlalchemy.Table(
                    'mssql_nvarchar_length_downsize',
                    metadata,
                    sqlalchemy.Column(
                        'id',
                        sqlalchemy.NVARCHAR(length=100),
                        primary_key=True
                    ),
                    sqlalchemy.Column(
                        'some_data', sqlalchemy.NVARCHAR(length=None)
                    ),
                )
            ).compile(writer.connection.engine)
            metadata.create_all(writer.connection.engine)

            writer.write_table(
                TableSpec(
                    **{
                        'name':
                            'mssql_nvarchar_length',
                        'headings': ['id', 'some_data'],
                        'rows': [
                            [
                                'bizzle', (b'\0' * 800).decode('utf-8'),
                                (b'\0' * 800).decode('utf-8')
                            ],
                            [
                                'bazzle', (b'\0' * 500).decode('utf-8'),
                                (b'\0' * 800).decode('utf-8')
                            ],
                        ]
                    }
                )
            )
            result = self._get_column_lengths(
                writer.connection, 'mssql_nvarchar_length_downsize'
            )
            assert result['some_data'] == ('some_data', 'nvarchar', -1)

    def test_big_lump_of_poo(self, writer):
        with writer:
            writer.write_table(
                TableSpec(
                    **{
                        'name': 'foo_with_emoji',
                        'headings': ['id', 'fun_to_be_had'],
                        'rows': [
                            ['A steaming poo', '💩'],
                            ['2020', '😷'],
                        ],
                    }
                )
            )

    def _get_column_lengths(self, connection, table_name):
        return {
            row['COLUMN_NAME']: row for row in connection.execute(
                "SELECT COLUMN_NAME, DATA_TYPE, CHARACTER_MAXIMUM_LENGTH "
                "FROM INFORMATION_SCHEMA.COLUMNS "
                "WHERE TABLE_NAME = '{}';".format(table_name)
            )
        }
